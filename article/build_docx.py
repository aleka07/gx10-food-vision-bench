"""Build the conference paper .docx and estimate its printed length.

Formatting follows the ATU conference requirements: A4, 2 cm margins, Times New
Roman 14 pt, single line spacing, 1 cm first-line indent, UDC in the top-left
corner, uppercase bold centred title, author block under it, numbered references
in square brackets.

Times New Roman itself is not installed on the build host, so the page estimate
uses Liberation Serif, which is metric-compatible with it. The estimate is a
sanity check for the 3-page limit, not a substitute for opening the file.
"""
from __future__ import annotations

import argparse
import math
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml.ns import qn

FONT = "Times New Roman"
BODY_PT = 14
TABLE_PT = 12
CAPTION_PT = 12
LINE_PT = 16.2          # single spacing for 14 pt Times New Roman in Word
PAGE_TEXT_HEIGHT_PT = 728.0   # A4 height minus 2 cm margins top and bottom
TEXT_WIDTH_CM = 17.0
PAGE_TEXT_WIDTH_PT = TEXT_WIDTH_CM * 28.3465
LIBERATION = "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf"


def _set_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(BODY_PT)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), FONT)
    rfonts.set(qn("w:cs"), FONT)
    pf = style.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    for section in doc.sections:
        section.page_height = Cm(29.7)
        section.page_width = Cm(21.0)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)


def para(
    doc: Document,
    text: str,
    *,
    size: int = BODY_PT,
    bold: bool = False,
    italic: bool = False,
    align: str = "justify",
    indent_cm: float = 1.0,
    space_before_pt: float = 0.0,
    space_after_pt: float = 0.0,
):
    p = doc.add_paragraph()
    p.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]
    p.paragraph_format.first_line_indent = Cm(indent_cm)
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return p


def add_table(doc: Document, caption: str, headers: list[str], rows: list[list[str]]) -> None:
    para(doc, caption, size=CAPTION_PT, align="left", indent_cm=0.0, space_before_pt=6)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        run.bold = True
        run.font.size = Pt(TABLE_PT)
        run.font.name = FONT
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            run.font.size = Pt(TABLE_PT)
            run.font.name = FONT
    para(doc, "", size=6, indent_cm=0.0)


def add_figure(doc: Document, path: Path, width_cm: float, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    para(doc, caption, size=CAPTION_PT, align="center", indent_cm=0.0)


def _wrap_lines(text: str, size_pt: float, width_pt: float, first_indent_pt: float = 0.0) -> int:
    """Estimate wrapped line count using a Times-metric font."""
    try:
        from PIL import ImageFont

        font = ImageFont.truetype(LIBERATION, int(round(size_pt * 96 / 72)))
        measure = lambda s: font.getlength(s) * 72 / 96  # noqa: E731
    except Exception:  # pragma: no cover - fallback keeps the estimate usable
        measure = lambda s: len(s) * size_pt * 0.5  # noqa: E731

    words = text.split()
    if not words:
        return 1
    lines, current, limit = 1, "", width_pt - first_indent_pt
    for word in words:
        candidate = f"{current} {word}".strip()
        if measure(candidate) <= limit or not current:
            current = candidate
        else:
            lines += 1
            current = word
            limit = width_pt
    return lines


def estimate_pages(content, figure_heights_cm: dict[str, float]) -> dict:
    """Rough page count for the assembled document."""
    height = 0.0
    detail = {}
    for text, size in (
        (content.udc, BODY_PT),
        (content.title, BODY_PT),
        (content.author_line, BODY_PT),
    ):
        lines = _wrap_lines(text, size, PAGE_TEXT_WIDTH_PT)
        height += lines * (LINE_PT if size >= BODY_PT else LINE_PT * size / BODY_PT)
        detail[text[:24]] = lines
    for text in content.paragraphs:
        lines = _wrap_lines(text, BODY_PT, PAGE_TEXT_WIDTH_PT, first_indent_pt=28.35)
        height += lines * LINE_PT
    for block in content.blocks:
        if block["kind"] == "table":
            height += (len(block["rows"]) + 1) * (TABLE_PT * 1.45 + 2) + CAPTION_PT * 1.5 + 6
        elif block["kind"] == "figure":
            caption_lines = _wrap_lines(block["caption"], CAPTION_PT, PAGE_TEXT_WIDTH_PT)
            height += figure_heights_cm[block["path"]] * 28.3465 + caption_lines * CAPTION_PT * 1.3 + 6
    height += _wrap_lines("Список литературы:", BODY_PT, PAGE_TEXT_WIDTH_PT) * LINE_PT
    for ref in content.references:
        height += _wrap_lines(ref, BODY_PT, PAGE_TEXT_WIDTH_PT, first_indent_pt=28.35) * LINE_PT
    pages = height / PAGE_TEXT_HEIGHT_PT
    return {"estimated_pages": round(pages, 2), "used_height_pt": round(height, 1), "detail": detail}


def build(content, out_path: Path, figures_dir: Path) -> dict:
    doc = Document()
    _set_style(doc)
    para(doc, content.udc, align="left", indent_cm=0.0)
    para(doc, content.title, bold=True, align="center", indent_cm=0.0)
    para(doc, content.author_line, align="center", indent_cm=0.0, space_after_pt=6)
    for text in content.paragraphs[: content.block_at] if hasattr(content, "block_at") else content.paragraphs:
        para(doc, text)

    figure_heights: dict[str, float] = {}
    for i, block in enumerate(content.blocks):
        if block["kind"] == "table":
            add_table(doc, block["caption"], block["headers"], block["rows"])
        elif block["kind"] == "figure":
            add_figure(doc, figures_dir / block["path"], block["width_cm"], block["caption"])
            figure_heights[block["path"]] = block.get("height_cm", block["width_cm"] * 0.42)
        if hasattr(content, "block_at") and i + 1 == content.block_at:
            for text in content.paragraphs[content.block_at :]:
                para(doc, text)
    if hasattr(content, "tail_at"):
        for index, text in enumerate(content.tail_paragraphs, start=content.tail_at):
            if index < len(content.paragraphs):
                continue
            para(doc, text)

    para(doc, "Список литературы:", indent_cm=0.0, space_before_pt=6)
    for number, ref in enumerate(content.references, start=1):
        para(doc, f"{number}. {ref}", indent_cm=1.0)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return estimate_pages(content, figure_heights)


if __name__ == "__main__":
    import importlib.util

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--content", default="article_text.py")
    parser.add_argument("--out", default="../Статья_МНПК_АТУ_2026_Амирханов.docx")
    parser.add_argument("--figures", default="../figures")
    args = parser.parse_args()

    spec = importlib.util.spec_from_file_location("article_text", Path(args.content).resolve())
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    stats = build(module, Path(args.out).resolve(), Path(args.figures).resolve())
    print(stats)
